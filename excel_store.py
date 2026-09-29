import os
from openpyxl import Workbook, load_workbook

# Always save next to this file, so the path problem from before can't happen
BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE_NAME = os.path.join(BASE_DIR, "attendance.xlsx")
FIXED_COLUMNS = 2  # Reg No, Name


def _open_workbook():
    if os.path.exists(FILE_NAME):
        return load_workbook(FILE_NAME)
    wb = Workbook()
    ws = wb.active
    ws.title = "Attendance"
    ws.append(["Reg No", "Name"])
    return wb


def _save(wb):
    try:
        wb.save(FILE_NAME)
        return True
    except PermissionError:
        print("Cannot save. Close attendance.xlsx in Excel and try again.")
        return False


def _clean(reg_no):
    return str(reg_no).strip().upper()


def get_students():
    """Return a dictionary {reg_no: name} of all registered students."""
    ws = _open_workbook().active
    students = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0]:
            students[_clean(row[0])] = row[1]
    return students


def add_student(reg_no, name):
    """Add a student. Returns False if the reg no already exists."""
    reg_no = _clean(reg_no)
    if reg_no in get_students():
        return False
    wb = _open_workbook()
    ws = wb.active
    row = [reg_no, name]
    row += ["Absent"] * (ws.max_column - FIXED_COLUMNS)  # absent for past classes
    ws.append(row)
    return _save(wb)


def start_class(label):
    """Add a new column for this class, with everyone set to Absent."""
    wb = _open_workbook()
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    if label in headers:
        return True  # class already started, keep existing marks
    col = ws.max_column + 1
    ws.cell(row=1, column=col, value=label)
    for r in range(2, ws.max_row + 1):
        ws.cell(row=r, column=col, value="Absent")
    return _save(wb)


def mark_present(reg_no, label):
    """Returns: marked, already, not_found, no_class or save_failed."""
    reg_no = _clean(reg_no)
    wb = _open_workbook()
    ws = wb.active
    headers = [cell.value for cell in ws[1]]
    if label not in headers:
        return "no_class"
    col = headers.index(label) + 1
    for r in range(2, ws.max_row + 1):
        if _clean(ws.cell(row=r, column=1).value) == reg_no:
            if ws.cell(row=r, column=col).value == "Present":
                return "already"
            ws.cell(row=r, column=col, value="Present")
            return "marked" if _save(wb) else "save_failed"
    return "not_found"

def delete_student(reg_no):
    reg_no = _clean(reg_no)
    wb = _open_workbook()
    ws = wb.active
    for r in range(2, ws.max_row + 1):
        if _clean(ws.cell(row=r, column=1).value) == reg_no:
            ws.delete_rows(r)
            return _save(wb)
    return False

if __name__ == "__main__":
    # Quick test: run this file directly
    add_student("26BCE0001", "Test One")
    add_student("26BCE0002", "Test Two")
    start_class("CSE 3:00 2026-09-28")
    print(mark_present("26bce0001", "CSE 3:00 2026-09-28"))
    print(mark_present("26bce0001", "CSE 3:00 2026-09-28"))
    print(mark_present("99XXX", "CSE 3:00 2026-09-28"))