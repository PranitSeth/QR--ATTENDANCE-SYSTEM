from excel_store import get_students, FILE_NAME, FIXED_COLUMNS
from openpyxl import load_workbook

def view_all_students():
    students = get_students()
    if not students:
        print("No students registered yet.")
        return
    print("\n--- All Registered Students ---")
    for reg_no, name in students.items():
        print(f"Reg No: {reg_no} | Name: {name}")

def _attendance_counts():
    wb = load_workbook(FILE_NAME)
    ws = wb.active
    counts = {}
    for row in ws.iter_rows(min_row=2, values_only=True):
        reg_no, name = row[0], row[1]
        if not reg_no:
            continue
        present_count = sum(1 for cell in row[FIXED_COLUMNS:] if cell == "Present")
        counts[reg_no] = {"name": name, "count": present_count}
    return counts

def show_attendance_summary():
    counts = _attendance_counts()
    if not counts:
        print("No students registered yet.")
        return
    print("\n--- Attendance Summary (classes present) ---")
    for reg_no, info in counts.items():
        print(f"Reg No: {reg_no} | Name: {info['name']} | Classes Present: {info['count']}")

def show_top_attendance(n=1):
    counts = _attendance_counts()
    if not counts:
        print("No students registered yet.")
        return
    ranked = sorted(counts.items(), key=lambda x: x[1]["count"], reverse=True)
    print(f"\n--- Top {n} Student(s) by Attendance ---")
    for reg_no, info in ranked[:n]:
        print(f"Reg No: {reg_no} | Name: {info['name']} | Classes Present: {info['count']}")

def show_kth_highest_attendance(k):
    counts = _attendance_counts()
    if not counts or k > len(counts) or k < 1:
        print("Invalid rank or no data available.")
        return
    ranked = sorted(counts.items(), key=lambda x: x[1]["count"], reverse=True)
    reg_no, info = ranked[k - 1]
    print(f"\n{k}-th highest attendance: Reg No {reg_no} | Name: {info['name']} | Classes Present: {info['count']}")